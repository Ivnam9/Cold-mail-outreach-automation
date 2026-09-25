# -*- coding: utf-8 -*-
"""Build a multi-sheet Excel workbook from generated emails — one sheet per domain
bucket, ordered largest-first, for human review before anything gets sent.

Usage:
    python build_workbook.py --input emails.json --output Professors.xlsx
"""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADERS = [
    "Name",
    "Email",
    "Designation",
    "Research Description",
    "University",
    "Timezone",
    "Subject",
    "Personalized Email",
    "Resume Path",
]

COLUMN_WIDTHS = {
    "A": 26,
    "B": 32,
    "C": 38,
    "D": 70,
    "E": 30,
    "F": 22,
    "G": 45,
    "H": 100,
    "I": 65,
}

HEADER_FILL = PatternFill(
    start_color="1F4E78",
    end_color="1F4E78",
    fill_type="solid",
)
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def build(records: list, out_path: str) -> None:
    groups: dict = {}
    for r in records:
        groups.setdefault(r.get("domain_bucket", "Other"), []).append(r)

    order = sorted(groups, key=lambda k: -len(groups[k]))

    wb = Workbook()
    wb.remove(wb.active)

    for bucket in order:
        rows = sorted(groups[bucket], key=lambda r: r.get("name", ""))
        ws = wb.create_sheet(title=bucket[:31])

        ws.append(HEADERS)

        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT

        for r in rows:
            ws.append([
                r.get("name", ""),
                r.get("email") or "",
                r.get("title", ""),
                r.get("bio", ""),
                r.get("university", ""),
                r.get("timezone", ""),
                r.get("subject", ""),
                r.get("email_draft", ""),
                r.get("resume_path", ""),
            ])

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = WRAP

        for col, width in COLUMN_WIDTHS.items():
            ws.column_dimensions[col].width = width

        ws.freeze_panes = "A2"

    wb.save(out_path)

    print(
        f"Saved {out_path} "
        f"({len(records)} people across {len(order)} sheets)"
    )

    for bucket in order:
        print(f"  {bucket}: {len(groups[bucket])}")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument(
        "--input",
        required=True,
        help="emails.json (output of email_generator.py)",
    )
    ap.add_argument(
        "--output",
        required=True,
        help="output .xlsx path",
    )
    args = ap.parse_args()

    records = json.loads(
        Path(args.input).read_text(encoding="utf-8")
    )

    build(records, args.output)


if __name__ == "__main__":
    main()